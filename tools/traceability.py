# Maps every test case ID in the sheet to the Playwright tests that carry it.
# Run from the project root:
#   python3 tools/traceability.py <sheet.xlsx> [report.xlsx]
import json, os, re, subprocess, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

args = sys.argv[1:]
results_files = []
while '--results' in args:  # optional, can be given more than once: JSON result files of runs
    k = args.index('--results')
    results_files.append(args[k + 1])
    del args[k:k + 2]
results_file = results_files[0] if results_files else None
final_file = None
if '--final' in args:  # optional: also write the presentation workbook
    k = args.index('--final')
    final_file = args[k + 1]
    del args[k:k + 2]
if not args:
    sys.exit('Usage: python3 tools/traceability.py <sheet.xlsx> [report.xlsx] [--results run.json ...] [--final workbook.xlsx]')
sheet_file = args[0]
report_file = args[1] if len(args) > 1 else 'traceability-report.xlsx'
ID = re.compile(r'ADM_\d{3}\.\d{3}')
SHORT = re.compile(r'(ADM_\d{3})\.\d{3}((?:\s*/\s*\.?\d{3}(?!\d))+)')

def ids_of(title):
    ids = ID.findall(title)
    for m in SHORT.finditer(title):  # shorthand such as ADM_002.004/005
        ids += [f"{m.group(1)}.{n}" for n in re.findall(r'\d{3}', m.group(2))]
    return list(dict.fromkeys(ids))

# 1. All tests, straight from Playwright (nothing is run)
json_file = os.environ.get('TRACE_JSON')  # only used to test this script itself
if not json_file:
    json_file = os.path.abspath('traceability-list.json')
    subprocess.run(['npx', 'playwright', 'test', '--list', '--reporter=json'],
                   env={**os.environ, 'PLAYWRIGHT_JSON_OUTPUT_NAME': json_file},
                   stdout=subprocess.DEVNULL, check=True)
data = json.load(open(json_file, encoding='utf8'))
cfg = data['config']
base = (cfg.get('projects') or [{}])[0].get('testDir') or cfg['rootDir']
found = {}
def walk(suite):
    for spec in suite.get('specs', []):
        f = os.path.relpath(os.path.join(base, spec['file']), os.getcwd()).replace(os.sep, '/')
        found.setdefault((f, spec['line'], spec['title']), {'title': spec['title'], 'file': f, 'line': spec['line']})
    for s in suite.get('suites', []):
        walk(s)
for s in data.get('suites', []):
    walk(s)
tests = list(found.values())

# 1b. Last run result per test (only when --results is given)
STATUS = {'expected': 'Pass', 'unexpected': 'Fail', 'flaky': 'Flaky', 'skipped': 'Skipped'}
RANK = ['Fail', 'Flaky', 'Pass', 'Skipped']  # worst first
run_status = {}
run_date = ''
for rf in results_files:
    rdata = json.load(open(rf, encoding='utf8'))
    run_date = max(run_date, (rdata.get('stats') or {}).get('startTime', '')[:10])
    def rwalk(suite):
        for spec in suite.get('specs', []):
            f = os.path.relpath(os.path.join(base, spec['file']), os.getcwd()).replace(os.sep, '/')
            sts = [STATUS.get(t.get('status'), 'Skipped') for t in spec.get('tests', [])]
            if sts:
                key = (f, spec['line'], spec['title'])
                new = min(sts, key=RANK.index)
                # the same test appears once per project: keep the worst result, never let Skipped hide a Pass
                run_status[key] = new if key not in run_status else min([run_status[key], new], key=RANK.index)
        for s2 in suite.get('suites', []):
            rwalk(s2)
    for s2 in rdata.get('suites', []):
        rwalk(s2)

# 2. Index tests by TC ID
by_id, no_id = {}, []
for t in tests:
    ids = ids_of(t['title'])
    if not ids:
        no_id.append(t)
    for i in ids:
        by_id.setdefault(i, []).append(t)

# 3. Read the sheet and compare
MUST = ('Automated', 'Partially Automated')
NONE = ('Not Automated', 'Better to Keep Manual', 'Not Suitable')
wb = openpyxl.load_workbook(sheet_file)
if 'TC Status' not in wb.sheetnames:
    sys.exit('The sheet file has no "TC Status" tab.')
rows = []
for r in wb['TC Status'].iter_rows(min_row=2, values_only=True):
    if r[0] and re.fullmatch(r'ADM_\d{3}\.\d{3}', str(r[0]).strip()):
        rows.append((str(r[0]).strip(), r[1], r[2], str(r[3] or '')))
sheet_ids = {r[0] for r in rows}

out = [['TC ID', 'Module', 'Description', 'Sheet status', 'Tests found', 'Files', 'Lines', 'Tags', 'Result', 'Last run'] ]
run_counts = {}
counts = {}
for tc, mod, desc, status in rows:
    t = by_id.get(tc, [])
    files = '\n'.join(dict.fromkeys(x['file'] for x in t))
    lines = '\n'.join(f"{os.path.basename(x['file'])}:{x['line']}" for x in t)
    tags = ' '.join(dict.fromkeys(g for x in t for g in re.findall(r'@\w+', x['title'])))
    if status in MUST and not t:
        res = 'MISSING: status says automated but no test carries this ID'
    elif status in NONE and t:
        res = 'CHECK: a test exists, the status may be outdated'
    elif status == 'Covered by other test' and not t:
        res = 'OK (covered by the form tests, no own ID)'
    else:
        res = 'OK'
    k = res.split(':')[0].split(' (')[0]
    counts[k] = counts.get(k, 0) + 1
    last = ''
    if results_file:
        sts = [run_status.get((x['file'], x['line'], x['title']), 'Skipped') for x in t]
        last = min(sts, key=RANK.index) if sts else 'No test'
        if last == 'Skipped' and any('/real-send/' in x['file'] or 'GuestLogs' in x['file'] for x in t):
            last = 'Skipped (opt-in real send)'
        run_counts[last] = run_counts.get(last, 0) + 1
    out.append([tc, mod, desc, status, len(t), files, lines, tags, res, last])

extra = sorted(i for i in by_id if i not in sheet_ids)

def put(name, table, widths, color_col=None):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for row in table:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    if color_col:
        for row in ws.iter_rows(min_row=2):
            v = str(row[color_col - 1].value)
            col = 'FBEBEB' if v.startswith('MISSING') else 'F9EEDA' if v.startswith('CHECK') else 'E5F4E4'
            row[color_col - 1].fill = PatternFill('solid', fgColor=col)
        if ws.max_column >= 10:
            for row in ws.iter_rows(min_row=2):
                v = str(row[9].value)
                col = {'Pass': 'E5F4E4', 'Fail': 'FBEBEB', 'Flaky': 'F9EEDA'}.get(v) or ('F0EFEC' if v.startswith('Skipped') else None)
                if col:
                    row[9].fill = PatternFill('solid', fgColor=col)

put('Traceability', out, [14, 24, 40, 22, 8, 50, 40, 18, 50, 12], 9)
put('Tests without a sheet ID', [['Test title', 'File', 'Line']] + [[t['title'], t['file'], t['line']] for t in no_id], [90, 50, 8])
put('IDs in tests, not in sheet', [['TC ID', 'Files']] + [[i, '\n'.join(dict.fromkeys(x['file'] for x in by_id[i]))] for i in extra], [16, 60])
wb.save(report_file)

# 4. Presentation workbook (optional)
if final_file:
    G, Y, R, B, Gr = 'E5F4E4', 'F9EEDA', 'FBEBEB', 'E7F1FB', 'F0EFEC'
    fill = {'Automated': G, 'Partially Automated': Y, 'Covered by other test': G, 'Not Automated': R, 'Better to Keep Manual': B, 'Not Suitable': Gr}
    notes = {r[0]: r[4] for r in wb['TC Status'].iter_rows(min_row=2, values_only=True) if r[0]}
    generic = {'Covered (109 of 138)', 'My Account done', 'Real-send cases need ALLOW_GUEST_SEND=1', None, ''}
    fb = openpyxl.load_workbook(sheet_file)
    orig = next(w for w in fb.worksheets if w.title not in ('TC Status', 'Status Summary'))
    orig.title = 'Original TC sheet'
    for n in list(fb.sheetnames):
        if n != 'Original TC sheet':
            del fb[n]
    def sheet(name, table, widths, idx):
        w = fb.create_sheet(name, idx)
        for r in table:
            w.append(r)
        for c in w[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='444444')
        for i, wd in enumerate(widths):
            w.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
        for row in w.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
        w.freeze_panes = 'A2'
        return w
    data = [['TC ID', 'Module', 'Description', 'Automation status', f'Last run ({run_date})', 'Test file(s)', 'Line(s)', 'Tags', 'Notes']]
    for tc, mod, desc, status, n, files, lines, tags, res, last in out[1:]:
        note = notes.get(tc)
        if note in generic: note = ''
        if note and note.startswith('tests/') and ' - ' in note: note = note.split(' - ', 1)[1]
        if status == 'Covered by other test': note = 'Same rules are tested in the form tests (ADM_022.022 to .032).'
        if last.startswith('Skipped'): note = 'Real send: runs only with ALLOW_GUEST_SEND=1 / ALLOW_MEMBER_SEND=1; not part of this run.'
        elif t_real := any('/real-send/' in f or 'GuestLogs' in f for f in files.split('\n')):
            note = 'Real send to the QA environment, run on demand with the switches on.'
        if last == 'No test': last = 'No automated test'
        data.append([tc, mod, desc, status, last, files, lines, tags, note])
    w = sheet('TC Status', data, [13, 24, 42, 22, 22, 48, 34, 16, 55], 0)
    w.auto_filter.ref = w.dimensions
    for row in w.iter_rows(min_row=2):
        row[3].fill = PatternFill('solid', fgColor=fill.get(row[3].value, Gr))
        v = str(row[4].value)
        c = {'Pass': G, 'Fail': R, 'Flaky': Y}.get(v) or (Gr if v.startswith(('Skipped', 'No')) else None)
        if c: row[4].fill = PatternFill('solid', fgColor=c)
    order = ['Automated', 'Partially Automated', 'Covered by other test', 'Not Automated', 'Better to Keep Manual', 'Not Suitable']
    by_mod = {}
    lastc, stc = {}, {}
    for r in data[1:]:
        m = r[0].split('.')[0]
        by_mod.setdefault(m, {}); by_mod[m][r[3]] = by_mod[m].get(r[3], 0) + 1
        stc[r[3]] = stc.get(r[3], 0) + 1
        lastc[r[4]] = lastc.get(r[4], 0) + 1
    S = [['Module'] + order + ['Total']]
    for m in sorted(by_mod):
        S.append([m] + [by_mod[m].get(x, 0) for x in order] + [sum(by_mod[m].values())])
    S.append(['All'] + [stc.get(x, 0) for x in order] + [len(data) - 1])
    S += [[], [f'Runs behind the Last run column ({run_date})']]
    for rf in results_files:
        st = json.load(open(rf, encoding='utf8')).get('stats') or {}
        S.append([os.path.basename(rf), f"{st.get('expected', 0) + st.get('unexpected', 0) + st.get('flaky', 0) + st.get('skipped', 0)} test runs", f"{st.get('expected', 0)} passed", f"{st.get('unexpected', 0)} failed, {st.get('flaky', 0)} flaky", f"{st.get('skipped', 0)} skipped"])
    S += [[], ['Last run per test case'], ['Pass', lastc.get('Pass', 0)], ['Fail / Flaky', lastc.get('Fail', 0) + lastc.get('Flaky', 0)],
          ['Skipped (not run in this report)', sum(v for k, v in lastc.items() if k.startswith('Skipped'))],
          ['No automated test (manual, blocked, parked, not suitable, covered by other test)', lastc.get('No automated test', 0)], [],
          [f"Every TC ID in the sheet was matched to a test title by script: {counts.get('OK', 0)} of {len(rows)} consistent with the code; tests referring to an ID outside the sheet: {len(extra)}."]]
    sm = sheet('Summary', S, [44, 20, 16, 22, 16, 18, 10], 0)
    for c in sm[len(by_mod) + 2]: c.font = Font(bold=True)
    how = [['How to find and run a test case', ''],
           ['Find the test for a TC ID', 'grep -rn "ADM_027.073" tests'],
           ['Run one test case (each run also runs 2 setup tests)', 'npx playwright test --project=chromium-super-admin --grep "ADM_027.073"'],
           ['Run one test file', 'npx playwright test tests/notifications/list/notificationList.spec.js --project=chromium-super-admin'],
           ['Notifications, full suite as General Admin', 'npm run test:notif:ga'],
           ['Notifications, critical set as Super Admin (15 tests)', 'npm run test:notif:sa'],
           ['Other modules', 'npm run test:auth | test:account | test:winners | test:users | test:tags'],
           ['Everything, all four projects', 'npm run test:all'],
           ['Real sends to the QA environment (opt-in)', 'npm run test:real-send:guest | npm run test:real-send:member'],
           ['Monthly full run and this workbook', 'bash tools/monthly-run.sh'], [],
           ['Status meanings', ''],
           ['Automated', 'A test carries this TC ID and runs in the suite.'],
           ['Partially Automated', 'Part is automated; the rest is checked by hand (for example the phone side of a real send).'],
           ['Covered by other test', 'No test of its own; the same rule is checked by another test.'],
           ['Not Automated', 'Blocked (needs a way to read email links), parked (needs a developer decision) or waiting for the Figma check.'],
           ['Better to Keep Manual', 'Needs a person or a device (visual match, real inbox, the app on a phone).'],
           ['Not Suitable', 'Cannot be automated in a useful way (for example a 24-hour expiry).']]
    hw = fb.create_sheet('How to run', 2)
    for r in how: hw.append(r)
    hw.column_dimensions['A'].width = 62; hw.column_dimensions['B'].width = 95
    for row in hw.iter_rows():
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
    for r in (1, 12): hw.cell(r, 1).font = Font(bold=True, size=12)
    sheet('Tests without a TC ID', [['Test title', 'File', 'Line']] + [[t['title'], t['file'], t['line']] for t in no_id], [90, 55, 8], 3)
    fb.active = 0
    fb.save(final_file)
    print(f'Presentation workbook written to {final_file}')

print(f"Tests found: {len(tests)} ({len(tests) - len(no_id)} with a TC ID, {len(no_id)} without)")
print(f"Sheet cases: {len(rows)}")
print('Result:', counts)
if results_file:
    print(f'Last run ({run_date}):', run_counts)
print(f"IDs in tests but not in the sheet: {len(extra)}" + (' -> ' + ', '.join(extra) if extra else ''))
for r in [r for r in out[1:] if not r[8].startswith('OK')][:40]:
    print(f"{r[0]}  [{r[3]}]  {r[8]}")
print(f"Report written to {report_file}")
